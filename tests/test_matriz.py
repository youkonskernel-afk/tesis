#!/usr/bin/env python3
"""scripts/matriz.py: la matriz de features por locus.

Lo que protege, por orden de gravedad:

  1. Que el duplicado NUNCA entre al entrenamiento. Es la validación
     independiente: si entra, deja de serlo, y nada lo delata.
  2. Que el etiquetado sea el de docs/positivos.md: semilla 5' exacta, holgura en
     3', sin corrimiento del 5'. Un positivo falso en PU es peor que un no
     etiquetado.
  3. Que una feature mal calculada no pase en silencio: el 5' de un read en la
     hebra - es su último nt, la ventana del pliegue va en la hebra del RNA
     mayoritario, la fracción multimapeada cuenta P y R.
  4. Que una fila sin pliegue (NaN) no llegue a loo_cv.
  5. Que tradeoff se llame con la ventana del proyecto y parado adentro de -o.

La lógica se prueba con reads de mentira. pysam y ViennaRNA solo se usan en dos
adaptadores finos; si están instalados, se prueban contra un BAM de verdad.
"""
import contextlib
import io
import math
import os
import pathlib
import random
import stat
import sys
import tempfile

RAIZ = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import matriz  # noqa: E402

FALLAS = 0


def chk(n, ok, extra=""):
    global FALLAS
    print(("  ok   " if ok else "  MAL  ") + n + (f"  [{extra}]" if not ok and extra else ""))
    if not ok:
        FALLAS += 1


def fila_loci(**k):
    base = dict(Locus="chr1:100-200", Name="locus_1", Length="100", Reads="0", RPM="10.0",
                UniqueReads="3", FracTop="1.0", Strand="+", MajorRNA="", MajorRNAReads="0",
                Complexity="0.03", Gap="5", skew="0.8", size_1n="(22,)", size_1n_depth="80",
                size_2n="(21, 22)", size_2n_depth="85", size_3n="(21, 22, 23)",
                size_3n_depth="90", sizecall="22")
    base.update({kk: str(v) for kk, v in k.items()})
    return base


def capturar(fn, *a, **k):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        try:
            r = fn(*a, **k)
            e = None
        except SystemExit as x:
            r, e = None, x
    return r, e, buf.getvalue()


MAY = "TGAGGTAGTAGGTTGTATAGTT"            # 22 nt, empieza con U
GENOMA = "".join(random.Random(3).choice("ACGT") for _ in range(400))
LARGOS = {"chr1": len(GENOMA)}


def genoma(chrom, ini, fin):
    return GENOMA[ini:fin]


print("== 1. features de los reads")
# 80 copias del mayoritario en + (110-132), 10 de 24 nt en + con otro 5' (P),
# 10 de 21 nt en - (R). 2 de 4 librerías.
reads = ([(MAY, "+", 110, 132, "L1", "U")] * 80
         + [("A" + "C" * 23, "+", 115, 139, "L2", "P")] * 10
         + [("G" * 21, "-", 150, 171, "L1", "R")] * 10)
plegados = []


def fold_falso(seq):
    plegados.append(seq)
    return ("(" * len(seq), -30.0)


filas = matriz.construir([fila_loci(Reads=100, MajorRNA=MAY.replace("T", "U"), FracTop=0.9)],
                         lambda c, i, f: iter(reads), 4, genoma, LARGOS, fold_falso, "gadmo", "primario")
f = filas[0]
chk("reads_bam = reads leídos (100)", f["reads_bam"] == 100, f["reads_bam"])
chk("y reads_yasma de loci.txt", f["reads_yasma"] == 100, f["reads_yasma"])
chk("locus_id lleva org y rol", f["locus_id"] == "gadmo_primario:locus_1", f["locus_id"])
chk("frac_22 = 0.8", math.isclose(f["frac_22"], 0.8), f["frac_22"])
chk("frac_24 = 0.1 y frac_21 = 0.1", math.isclose(f["frac_24"], 0.1) and math.isclose(f["frac_21"], 0.1))
chk("frac_5U = 0.8 (el mayoritario empieza con U)", math.isclose(f["frac_5U"], 0.8), f["frac_5U"])
chk("frac_5A = 0.1, frac_5G = 0.1", math.isclose(f["frac_5A"], 0.1) and math.isclose(f["frac_5G"], 0.1))
chk("precision_5p: 80 de 90 en la hebra dominante", math.isclose(f["precision_5p"], 80 / 90, rel_tol=1e-4),
    f["precision_5p"])
chk("frac_multimapeo cuenta P y R (20 de 100)", math.isclose(f["frac_multimapeo"], 0.2), f["frac_multimapeo"])
chk("frac_librerias = 2 de 4", math.isclose(f["frac_librerias"], 0.5), f["frac_librerias"])
chk("sesgo_hebra = |2·0.9-1| = 0.8", math.isclose(f["sesgo_hebra"], 0.8), f["sesgo_hebra"])
chk("largo, GC y 5'U del mayoritario",
    f["largo_mayoritario"] == 22 and f["mayoritario_5U"] == 1
    and math.isclose(f["gc_mayoritario"], (MAY.count("G") + MAY.count("C")) / 22, rel_tol=1e-4))
chk("log_reads = log10(101)", math.isclose(f["log_reads"], math.log10(101), rel_tol=1e-5))
chk("todas las FEATURES están", all(c in f for c in matriz.FEATURES),
    [c for c in matriz.FEATURES if c not in f])

print("== 2. el 5' de la hebra - es su último nt")
# 3 largos distintos que terminan en el mismo nt del genoma: en la hebra - el
# 5' es ese último nt, así que la precisión es 1.0. Mirando el inicio daría 1/3.
reads_m = [("T" * 22, "-", 200, 222, "L1", "U"), ("T" * 21, "-", 201, 222, "L1", "U"),
           ("T" * 23, "-", 199, 222, "L1", "U")] * 10
f = matriz.construir([fila_loci(Locus="chr1:190-230", Reads=30, MajorRNA="U" * 22, FracTop=0.0)],
                     lambda c, i, fn: iter(reads_m), 1, genoma, LARGOS, fold_falso, "gadmo", "primario")[0]
chk("precision_5p = 1.0", math.isclose(f["precision_5p"], 1.0), f["precision_5p"])

print("== 3. el pliegue va en la hebra del RNA mayoritario, centrado en él")
plegados.clear()
apareado = {}


def fold_marca(seq):
    # Apareado SOLO donde tiene que estar el mayoritario: 80 nt de flanco a la
    # izquierda en la hebra orientada. Si la ventana o el offset se corren, la
    # fracción apareada baja de 1.
    plegados.append(seq)
    est = "." * 80 + "(" * 22 + "." * (len(seq) - 102)
    return (est, -40.0)


# en +: el mayoritario en 150-172; la ventana es 70-252
r_mas = [(GENOMA[150:172], "+", 150, 172, "L1", "U")] * 5
f = matriz.construir([fila_loci(Locus="chr1:140-180", Reads=5, MajorRNA=GENOMA[150:172])],
                     lambda c, i, fn: iter(r_mas), 1, genoma, LARGOS, fold_marca, "gadmo", "primario")[0]
chk("en +: pliega la ventana ±80", plegados[-1] == GENOMA[70:252].replace("T", "U"), len(plegados[-1]))
chk("en +: el mayoritario cae apareado", math.isclose(f["frac_pareado_mayoritario"], 1.0),
    f["frac_pareado_mayoritario"])
chk("mfe y mfe_por_nt", f["mfe"] == -40.0 and math.isclose(f["mfe_por_nt"], -40 / 182, rel_tol=1e-3))
# en -: la molécula es el complemento reverso de 150-172
rc = matriz.revcomp(GENOMA[150:172])
r_menos = [(rc, "-", 150, 172, "L1", "U")] * 5
f = matriz.construir([fila_loci(Locus="chr1:140-180", Reads=5, MajorRNA=rc, FracTop=0.0)],
                     lambda c, i, fn: iter(r_menos), 1, genoma, LARGOS, fold_marca, "gadmo", "primario")[0]
chk("en -: pliega el complemento reverso de la ventana",
    plegados[-1] == matriz.revcomp(GENOMA[70:252]).replace("T", "U"))
chk("en -: el mayoritario cae apareado (offset desde el otro lado)",
    math.isclose(f["frac_pareado_mayoritario"], 1.0), f["frac_pareado_mayoritario"])
# en - y pegado al borde: la ventana queda asimetrica (0-132), y el offset del
# mayoritario en la hebra orientada es 132-52 = 80, no 30. Centrado en la
# ventana, los dos offsets dan igual y un error en el de - no se ve.
rc_b = matriz.revcomp(GENOMA[30:52])
r_menos_b = [(rc_b, "-", 30, 52, "L1", "U")] * 5
f = matriz.construir([fila_loci(Locus="chr1:20-60", Reads=5, MajorRNA=rc_b, FracTop=0.0)],
                     lambda c, i, fn: iter(r_menos_b), 1, genoma, LARGOS, fold_marca, "gadmo", "primario")[0]
chk("en - junto al borde: ventana recortada y en complemento reverso",
    plegados[-1] == matriz.revcomp(GENOMA[0:132]).replace("T", "U"))
chk("en - junto al borde: el offset se cuenta desde el final de la ventana",
    math.isclose(f["frac_pareado_mayoritario"], 1.0), f["frac_pareado_mayoritario"])
# en el borde del cromosoma la ventana se recorta, no revienta
r_borde = [(GENOMA[10:32], "+", 10, 32, "L1", "U")] * 5
f = matriz.construir([fila_loci(Locus="chr1:0-40", Reads=5, MajorRNA=GENOMA[10:32])],
                     lambda c, i, fn: iter(r_borde), 1, genoma, LARGOS, fold_falso, "gadmo", "primario")[0]
chk("en el borde: ventana recortada a 0", plegados[-1] == GENOMA[0:112].replace("T", "U"))

print("== 4. sin el mayoritario entre los reads, el pliegue es NaN")
f = matriz.construir([fila_loci(Reads=5, MajorRNA="A" * 22)],
                     lambda c, i, fn: iter(r_mas), 1, genoma, LARGOS, fold_falso, "gadmo", "primario")[0]
chk("mfe NaN", math.isnan(f["mfe"]), f["mfe"])

print("== 5. leer loci.txt de YASMA")
with tempfile.TemporaryDirectory() as d:
    d = pathlib.Path(d)
    (d / "tradeoff").mkdir()
    loci = d / "tradeoff" / "loci.txt"
    loci.write_text("\t".join(matriz.CABECERA_LOCI) + "\n"
                    + "\t".join(fila_loci(Reads=100, MajorRNA="UGA").values()) + "\n"
                    + "\t".join(fila_loci(Name="locus_2", Reads=0, MajorRNA="NA").values()) + "\n")
    chk("encuentra tradeoff/loci.txt", matriz.encontrar_loci(d) == loci)
    chk("saltea el locus sin reads", len(matriz.leer_loci(loci)) == 1)
    (d / "tradeoff_otro").mkdir()
    (d / "tradeoff_otro" / "loci.txt").write_text("x")
    _, e, _ = capturar(matriz.encontrar_loci, d)
    chk("con dos anotaciones no adivina", e is not None and "hay 2" in str(e.code), repr(e))
    otra = d / "malo.txt"
    otra.write_text("Locus\tName\tReads\nchr1:1-2\tl\t5\n")
    _, e, _ = capturar(matriz.leer_loci, otra)
    chk("una cabecera que no es la de v1.1.1 corta", e is not None and "v1.1.1" in str(e.code), repr(e))

print("== 6. anotar: tradeoff con la ventana del proyecto, parado adentro de -o")
with tempfile.TemporaryDirectory() as d:
    d = pathlib.Path(d)
    bam = d / "x.bam"
    bam.write_bytes(b"BAM")
    falso = d / "yasma"
    CAB_TXT = "\t".join(matriz.CABECERA_LOCI)
    FILA_TXT = "\t".join(fila_loci(Reads=100, MajorRNA="UGA").values())
    falso.write_text(f"""#!/usr/bin/env python3
import os, sys, pathlib
open({str(d / 'args')!r}, 'w').write(' '.join(sys.argv[1:]) + '\\n' + os.getcwd())
o = pathlib.Path(sys.argv[sys.argv.index('-o') + 1])
(o / 'tradeoff').mkdir(exist_ok=True)
filas = int(os.environ.get('FILAS', '1'))
cab = {CAB_TXT!r}
fila = {FILA_TXT!r}
(o / 'tradeoff' / 'loci.txt').write_text(cab + '\\n' + (fila + '\\n') * filas)
rc = int(os.environ.get('RC', '0'))
if rc < 0:
    import signal
    os.kill(os.getpid(), -rc)
sys.exit(rc)
""")
    falso.chmod(falso.stat().st_mode | stat.S_IEXEC)
    out = d / "anot"
    SI, NO = ["true"], ["false"]
    _, e, _ = capturar(matriz.anotar, bam, d / "anot0", str(falso), NO)
    chk("sin el parche de memoria no anota", e is not None and "tradeoff_parche.py" in str(e.code)
        and not (d / "args").exists(), repr(e))
    r, e, txt = capturar(matriz.anotar, bam, out, str(falso), SI)
    args, cwd = (d / "args").read_text().splitlines()
    chk("llama a tradeoff", args.startswith("tradeoff"), args)
    chk("con -a absoluto", f"-a {bam.resolve()}" in args, args)
    chk("con -o absoluto", f"-o {out.resolve()}" in args, args)
    chk("con la ventana 15-50 del proyecto", "--min_length 15 --max_length 50" in args, args)
    chk("y los filtros explícitos", "--min_abundance 50" in args and "--max_skew 0.9" in args, args)
    chk("parado adentro de -o", cwd == str(out.resolve()), cwd)
    chk("devuelve el loci.txt", r == out.resolve() / "tradeoff" / "loci.txt", r)
    chk("informa la memoria pico", "memoria pico de tradeoff:" in txt, txt[-300:])
    os.environ["RC"] = "2"
    _, e, _ = capturar(matriz.anotar, bam, d / "anot2", str(falso), SI)
    chk("si yasma falla, corta", e is not None and "salió con 2" in str(e.code), repr(e))
    os.environ["RC"] = "-9"
    _, e, _ = capturar(matriz.anotar, bam, d / "anot4", str(falso), SI)
    chk("si lo mata el sistema (-9), dice que es memoria", e is not None
        and "sin memoria" in str(e.code), repr(e))
    os.environ["RC"], os.environ["FILAS"] = "0", "0"
    _, e, _ = capturar(matriz.anotar, bam, d / "anot3", str(falso), SI)
    chk("si no anotó ningún locus, corta", e is not None and "ningún locus" in str(e.code), repr(e))
    del os.environ["RC"], os.environ["FILAS"]

print("== 7. etiquetar: semilla exacta, holgura en 3', sin correr el 5'")
MIR = "UGGAAUGUAAAGAAGUAUGUAU"  # 22
maduros = [("gmo-Mir-1", matriz.adn(MIR)), ("gmo-Mir-1-b", matriz.adn(MIR[:12] + "CC" + MIR[14:]))]


def y_de(seq, **k):
    fila = {"mayoritario": seq}
    matriz.etiquetar([fila], maduros, **k)
    return fila


chk("idéntico: positivo", y_de(MIR)["y"] == 1 and y_de(MIR)["y_mm"] == 0)
chk("con T en vez de U, igual", y_de(MIR.replace("U", "T"))["y"] == 1)
chk("+2 nt en 3': positivo, dlen 2", y_de(MIR + "AA")["y"] == 1 and y_de(MIR + "AA")["y_dlen"] == 2)
chk("-3 nt en 3': positivo", y_de(MIR[:-3])["y"] == 1)
f4 = y_de(MIR + "AAAA")
chk("+4 nt en 3': NO, pero queda la referencia", f4["y"] == 0 and f4["y_ref"] == "gmo-Mir-1", f4)
m1 = MIR[:15] + ("A" if MIR[15] != "A" else "C") + MIR[16:]
chk("1 diferencia en el cuerpo: positivo", y_de(m1)["y"] == 1 and y_de(m1)["y_mm"] == 1)
m2 = m1[:18] + ("A" if m1[18] != "A" else "C") + m1[19:]
f2 = y_de(m2)
chk("2 diferencias: NO, mm=2", f2["y"] == 0 and f2["y_mm"] == 2, f2)
chk("2 diferencias con --max-mm 2: sí", y_de(m2, max_mm=2)["y"] == 1)
ms = MIR[:3] + ("C" if MIR[3] != "C" else "G") + MIR[4:]
chk("1 diferencia en la semilla: sin referencia", y_de(ms)["y"] == 0 and y_de(ms)["y_ref"] == "", y_de(ms))
chk("5' corrido una base: sin referencia", y_de(MIR[1:])["y_ref"] == "", y_de(MIR[1:]))
chk("de la familia, elige el de menos diferencias",
    y_de(MIR[:12] + "CC" + MIR[14:])["y_ref"] == "gmo-Mir-1-b")

print("== 8. exportar: solo primarios de entrenamiento, nunca el duplicado")
with tempfile.TemporaryDirectory() as d:
    d = pathlib.Path(d)
    base = {c: "0.5" for c in matriz.FEATURES}
    cols = matriz.META + matriz.FEATURES + matriz.ETIQUETA

    def matriz_de(nombre, org, rol, ys, **extra):
        filas = [dict(base, org=org, rol=rol, locus_id=f"{org}_{rol}:l{i}", y=y, y_ref="r" if y else "",
                      y_mm=0, y_dlen=0, reads_yasma=1, reads_bam=1, **extra) for i, y in enumerate(ys)]
        matriz.escribir(filas, d / nombre, cols)
        return d / nombre

    mp = matriz_de("gp.tsv", "gadmo", "primario", [1, 0, 0])
    md = matriz_de("gd.tsv", "gadmo", "duplicado", [1, 1])
    ma = matriz_de("pp.tsv", "prupe", "primario", [0, 1])
    filas, afuera = matriz.exportar([mp, md, ma])
    chk("solo las 3 de gadmo primario", len(filas) == 3 and {f["locus_id"].split(":")[0] for f in filas}
        == {"gadmo_primario"}, [f["locus_id"] for f in filas])
    chk("el duplicado afuera, y dice por qué",
        afuera[("gadmo", "duplicado", "duplicado: es la validación independiente")] == 2, afuera)
    chk("la aplicación afuera, y dice por qué",
        afuera[("prupe", "primario", "aplicación: se predice, no se entrena")] == 2, afuera)
    out = d / "features.tsv"
    r, e, txt = capturar(matriz.main, ["exportar", str(mp), str(md), str(ma), "-o", str(out)])
    cab = out.read_text().splitlines()[0].split("\t")
    chk("columnas: org, locus_id, y y las FEATURES, nada más", cab == ["org", "locus_id", "y"] + matriz.FEATURES,
        cab)
    chk("sin columnas que filtren la etiqueta (y_ref, y_mm...)", not any(c.startswith("y_") for c in cab))
    chk("avisa que con un organismo no mide transferencia", "menos de 3 organismos" in txt, txt)
    # sin etiquetar
    sin_y = d / "sin_y.tsv"
    matriz.escribir([dict(base, org="gadmo", rol="primario", locus_id="x")], sin_y, matriz.META + matriz.FEATURES)
    _, e, _ = capturar(matriz.exportar, [sin_y])
    chk("una matriz sin etiquetar corta y manda a etiquetar", e is not None and "etiquetar" in str(e.code),
        repr(e))
    # NaN en una feature
    con_nan = matriz_de("nan.tsv", "gadmo", "primario", [1], )
    txt_nan = con_nan.read_text().replace("\t0.5\t", "\tnan\t", 1)
    con_nan.write_text(txt_nan)
    _, e, _ = capturar(matriz.exportar, [con_nan])
    chk("una feature NaN corta, nombrando el locus", e is not None and "NaN" in str(e.code), repr(e))
    _, e, _ = capturar(matriz.exportar, [mp, mp])
    chk("la misma matriz dos veces corta", e is not None and "repetido" in str(e.code), repr(e))

print("== 8c. excel: una hoja por proyecto, con diccionario y resumen")
try:
    import openpyxl
except ImportError:
    openpyxl = None
if openpyxl is None:
    print("  (salteado: sin openpyxl)")
else:
    with tempfile.TemporaryDirectory() as d:
        d = pathlib.Path(d)
        base = {c: "0.5" for c in matriz.FEATURES}
        con_y = matriz.META + matriz.FEATURES + matriz.ETIQUETA
        filas = [dict(base, org="gadmo", rol="primario", locus_id=f"gadmo_primario:l{i}",
                      locus="c1:1-50", mayoritario="UGA", y=y, y_ref="r" if y else "", y_mm=0,
                      y_dlen=0, reads_yasma=10, reads_bam=10) for i, y in enumerate([1, 0, 1])]
        filas[1]["mfe"] = "nan"
        matriz.escribir(filas, d / "gp.tsv", con_y)
        sin_y = [dict(base, org="sclsc", rol="duplicado", locus_id="sclsc_duplicado:l0",
                      reads_yasma=5, reads_bam=5)]
        matriz.escribir(sin_y, d / "sd.tsv", matriz.META + matriz.FEATURES)
        out = d / "f.xlsx"
        r, e, txt = capturar(matriz.main, ["excel", str(d / "gp.tsv"), str(d / "sd.tsv"), "-o", str(out)])
        wb = openpyxl.load_workbook(out)
        chk("hojas: Resumen, Diccionario y una por proyecto",
            wb.sheetnames == ["Resumen", "Diccionario", "gadmo_primario", "sclsc_duplicado"], wb.sheetnames)
        ws = wb["gadmo_primario"]
        cab = [c.value for c in ws[1]]
        chk("columnas: identificación, etiqueta y las 31 features, en ese orden",
            cab == matriz.META + matriz.ETIQUETA + matriz.FEATURES, cab[:12])
        fila2 = {h: c.value for h, c in zip(cab, ws[2])}
        chk("los números son números, no texto",
            isinstance(fila2["log_largo"], float) and fila2["reads_bam"] == 10, fila2["log_largo"])
        chk("un NaN queda vacío, no 'nan'", {h: c.value for h, c in zip(cab, ws[3])}["mfe"] is None)
        chk("sin etiqueta no inventa la columna y",
            "y" not in [c.value for c in wb["sclsc_duplicado"][1]])
        res = wb["Resumen"]
        chk("el resumen cuenta los loci con fórmula sobre la hoja",
            res["E2"].value == "=COUNTA('gadmo_primario'!C:C)-1", res["E2"].value)
        col_y = matriz._letra(cab.index("y") + 1)
        chk("y los positivos sobre la columna y",
            res["F2"].value == f"=COUNTIF('gadmo_primario'!{col_y}:{col_y},1)", res["F2"].value)
        chk("sin etiqueta lo dice", res["F3"].value == "sin etiquetar", res["F3"].value)
        chk("con el set_modelo de organismos.tsv",
            res["D2"].value == "entrenamiento" and res["D3"].value == "aplicacion")
        dic = {r[1]: r[2] for r in wb["Diccionario"].iter_rows(min_row=2, values_only=True)}
        chk("el diccionario explica cada feature", all(dic.get(f) for f in matriz.FEATURES),
            [f for f in matriz.FEATURES if not dic.get(f)])
        chk("y dice que y = 0 no es negativo", "no negativo" in (dic.get("y") or ""))
        chk("letra Arial", ws["K2"].font.name == "Arial")
        chk("recalcula al abrir", wb.calculation.fullCalcOnLoad is True)
        _, e, _ = capturar(matriz.excel, [d / "gp.tsv", d / "gp.tsv"], d / "g2.xlsx")
        chk("dos veces el mismo proyecto corta", e is not None and "dos matrices" in str(e.code), repr(e))
        (d / "malo.tsv").write_text("a\tb\n1\t2\n")
        _, e, _ = capturar(matriz.excel, [d / "malo.tsv"], d / "g3.xlsx")
        chk("algo que no es una matriz corta", e is not None and "no es una matriz" in str(e.code), repr(e))

print("== 8b. lo que usa el notebook: estado en Drive, maduros, resumen")
with tempfile.TemporaryDirectory() as d:
    d = pathlib.Path(d)
    for org, rol in (("gadmo", "primario"), ("gadmo", "duplicado"), ("prupe", "primario")):
        (d / "10_bam" / org).mkdir(parents=True, exist_ok=True)
        (d / "10_bam" / org / f"{rol}.bam").write_bytes(b"x" * 1000)
        (d / "10_bam" / org / f"{rol}.bam.bai").write_bytes(b"x")
    (d / "20_yasma" / "gadmo" / "primario" / "tradeoff").mkdir(parents=True)
    (d / "20_yasma" / "gadmo" / "primario" / "tradeoff" / "loci.txt").write_text("x")
    (d / "40_features" / "gadmo").mkdir(parents=True)
    (d / "40_features" / "gadmo" / "primario.matriz.tsv").write_text("x")
    est = {(f["org"], f["rol"]): f for f in matriz.estado_drive(d)}
    chk("ve los 3 BAM, y no el .bai como otro proyecto", sorted(est) ==
        [("gadmo", "duplicado"), ("gadmo", "primario"), ("prupe", "primario")], sorted(est))
    chk("gadmo/primario: anotación y matriz", est[("gadmo", "primario")]["anotacion"]
        and est[("gadmo", "primario")]["matriz"])
    chk("gadmo/duplicado: ninguna de las dos", not est[("gadmo", "duplicado")]["anotacion"]
        and not est[("gadmo", "duplicado")]["matriz"])
    chk("con su set_modelo", est[("gadmo", "primario")]["set_modelo"] == "entrenamiento"
        and est[("prupe", "primario")]["set_modelo"] == "aplicacion", est)
fa = "".join(f">gmo-Mir-{i}\n{MIR}\n" for i in range(25))
chk("un FASTA de maduros pasa", matriz.validar_maduros(fa) == (25, None), matriz.validar_maduros(fa))
n, mot = matriz.validar_maduros("<!DOCTYPE html><html>no encontrado</html>")
chk("una página HTML no pasa", mot is not None and "no es un FASTA" in mot, mot)
n, mot = matriz.validar_maduros(fa + ">pre\n" + "A" * 70 + "\n")
chk("precursores mezclados no pasan", mot is not None and "precursores" in mot, mot)
n, mot = matriz.validar_maduros(">a\n" + MIR + "\n")
chk("un FASTA de una secuencia no pasa", mot is not None and "solo 1" in mot, mot)
with tempfile.TemporaryDirectory() as d:
    d = pathlib.Path(d)
    m = d / "m.tsv"
    matriz.escribir([dict({c: "0.5" for c in matriz.FEATURES}, org="gadmo", rol="primario",
                          locus_id=f"l{i}", y=str(i % 2)) for i in range(4)], m,
                    matriz.META + matriz.FEATURES + matriz.ETIQUETA)
    r = matriz.resumen_matriz(m)
    chk("resumen: loci, positivos, features", (r["loci"], r["positivos"], r["features"])
        == (4, 2, len(matriz.FEATURES)), r)
    chk("y el sha256 del fichero", len(r["sha256"]) == 64, r)

print("== 9. con pysam y ViennaRNA: los adaptadores contra un BAM de verdad")
try:
    import pysam  # noqa: F401
    import RNA  # noqa: F401
    HAY = True
except ImportError:
    HAY = False
    print("  AVISO: sin pysam/ViennaRNA instalados, los adaptadores no se prueban acá "
          "(en Colab sí corren). pip install pysam ViennaRNA para probarlos.")

if HAY:
    import pysam
    with tempfile.TemporaryDirectory() as d:
        d = pathlib.Path(d)
        rnd = random.Random(11)
        g = "".join(rnd.choice("ACGT") for _ in range(600))
        (d / "g.fna").write_text(">chr1\n" + g + "\n")
        hdr = {"HD": {"VN": "1.6", "SO": "coordinate"}, "SQ": [{"SN": "chr1", "LN": len(g)}],
               "RG": [{"ID": "L1"}, {"ID": "L2"}, {"ID": "L3"}]}
        lecturas = []
        # el mayoritario en la hebra -: 300-322 del genoma
        for i in range(30):
            lecturas.append((300, g[300:322], True, "L1" if i % 2 else "L2", "U"))
        for i in range(10):
            lecturas.append((305, g[305:329], False, "L1", "P"))
        lecturas.sort(key=lambda x: x[0])
        sin_ordenar = d / "u.bam"
        with pysam.AlignmentFile(str(sin_ordenar), "wb", header=hdr) as out:
            for k, (pos, seq, rev, rg, xy) in enumerate(lecturas):
                a = pysam.AlignedSegment()
                a.query_name, a.query_sequence, a.flag = f"r{k}", seq, 16 if rev else 0
                a.reference_id, a.reference_start, a.mapping_quality = 0, pos, 255
                a.cigarstring, a.query_qualities = f"{len(seq)}M", pysam.qualitystring_to_array("I" * len(seq))
                a.set_tag("RG", rg)
                a.set_tag("XY", xy)
                out.write(a)
            # un no mapeado que cae en la región no cuenta (regla de YASMA)
            u = pysam.AlignedSegment()
            u.query_name, u.query_sequence, u.flag = "nm", "ACGT" * 5, 4
            u.reference_id, u.reference_start = 0, 310
            u.query_qualities = pysam.qualitystring_to_array("I" * 20)
            u.set_tag("RG", "L3")
            u.set_tag("XY", "N")
            out.write(u)
        bam = d / "x.bam"
        pysam.sort("-o", str(bam), str(sin_ordenar))
        pysam.index(str(bam))
        may = matriz.revcomp(g[300:322]).replace("T", "U")
        (d / "anot" / "tradeoff").mkdir(parents=True)
        (d / "anot" / "tradeoff" / "loci.txt").write_text(
            "\t".join(matriz.CABECERA_LOCI) + "\n"
            + "\t".join(fila_loci(Locus="chr1:290-340", Reads=40, MajorRNA=may, FracTop=0.25).values()) + "\n")
        salida = d / "m.tsv"
        rc, e, txt = capturar(matriz.main, ["features", str(d / "anot"), str(bam), "--org", "gadmo", "--rol",
                                            "primario", "--genoma", str(d / "g.fna"), "-o", str(salida)])
        chk("features sale 0", rc == 0 and e is None, (rc, e, txt[-300:]))
        _, filas = matriz.leer_tsv(salida)
        f = filas[0] if filas else {}
        chk("reads_bam = 40: el no mapeado no cuenta", f.get("reads_bam") == "40", f.get("reads_bam"))
        chk("y coincide con loci.txt, sin aviso", "AVISO" not in txt, txt[-300:])
        chk("encuentra el mayoritario en la hebra - (pliegue con número)",
            f.get("mfe") not in (None, "", "nan") and float(f["mfe"]) <= 0, f.get("mfe"))
        chk("frac_librerias = 2 de 3 (L3 solo tiene un no mapeado)",
            math.isclose(float(f.get("frac_librerias", -1)), 2 / 3, rel_tol=1e-4), f.get("frac_librerias"))
        chk("frac_multimapeo = 10 de 40", math.isclose(float(f.get("frac_multimapeo", -1)), 0.25),
            f.get("frac_multimapeo"))
        # el genoma equivocado se rechaza antes de calcular nada
        (d / "otro.fna").write_text(">chrX\n" + g + "\n")
        rc, e, _ = capturar(matriz.main, ["features", str(d / "anot"), str(bam), "--org", "gadmo", "--rol",
                                          "primario", "--genoma", str(d / "otro.fna"), "-o", str(salida)])
        chk("un genoma sin las secuencias del BAM corta", e is not None and "no es el FASTA" in str(e.code),
            repr(e))

print()
if FALLAS:
    print(f"{FALLAS} fallas")
    sys.exit(1)
print("TODO OK")
