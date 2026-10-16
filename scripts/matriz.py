#!/usr/bin/env python3
"""La matriz de features por locus que entra al PU learning.

    ./scripts/matriz.py anotar    <bam> <dir>                       yasma tradeoff sobre NUESTRO BAM
    ./scripts/matriz.py features  <dir> <bam> --org O --rol R --genoma G.fna -o M.tsv
    ./scripts/matriz.py etiquetar <M.tsv> <maduros.fa>... [-o M.tsv]
    ./scripts/matriz.py exportar  <M.tsv>... -o features.tsv        (lo que lee loo_cv.py)
    ./scripts/matriz.py excel     <M.tsv>... -o features.xlsx       para mirarlas, con diccionario

Una fila por locus, una columna por feature. El locus es lo que anota
`yasma tradeoff` sobre el BAM de `yasma align`: el alineamiento con `-m 50` es
el que define los loci, igual que en ShortStack.

QUÉ SALE DE DÓNDE, medido contra el código de YASMA v1.1.1 y no leído de la
documentación (ver docs/matriz.md):

  - `tradeoff` escribe en <dir>/tradeoff/: loci.txt (una fila por locus,
    columnas de assessClass en generics.py), reads.txt, loci.gff3, stats.txt.
  - Su `--max_length` NO filtra reads: el filtro por largo de samtools_view está
    comentado, y la llamada de tamaño (sizecall) usa 15-30 fijo en
    sizeClass.update(). Igual se pasa `--min_length 15 --max_length 50`, porque
    tradeoff lo escribe en inputs.json y params.json, y ahí tiene que quedar la
    ventana del proyecto y no el 30 por defecto.
  - Cuenta como reads de un locus todo read MAPEADO que solape el intervalo
    (`bamf.fetch(chrom, start, stop)`, coordenadas base 0). Acá se usa la misma
    regla, así que `reads_bam` tiene que dar igual a `Reads` de loci.txt: es el
    chequeo de que se está mirando lo mismo que YASMA miró.
  - La secuencia de cada read es `get_forward_sequence()`, o sea la molécula
    como se secuenció, no la hebra del genoma. El RNA mayoritario viene así, y
    así se compara contra los maduros.

LAS FEATURES (lista exacta en FEATURES; docs/matriz.md las justifica una por una):

  de loci.txt     largo, reads, RPM, lecturas únicas, complejidad, skew, sesgo
                  de hebra.
  de los reads    fracción por largo (15-17, 18 ... 25, 26-30, 31-50), fracción
                  por nucleótido 5', precisión del extremo 5' (qué fracción
                  empieza en la posición más común: Dicer corta preciso, la
                  degradación no), fracción multimapeada (marcas XY:Z:P/R de
                  yasma align) y fracción de librerías donde aparece el locus.
  del RNA mayor.  largo, GC, si empieza con U.
  del genoma      pliegue (ViennaRNA) de una ventana de ±FLANCO nt alrededor
                  del RNA mayoritario, en su hebra: MFE, MFE por nt, y qué
                  fracción del mayoritario queda apareada — un miRNA vive en el
                  tallo de una horquilla.

LO QUE NO ES FEATURE Y NO PUEDE SERLO: si el locus reaparece en el duplicado.
El duplicado es la validación independiente; meterlo como señal la anula.

ETIQUETADO (docs/positivos.md): por secuencia, no por coordenada. Un locus es
positivo si su RNA mayoritario coincide con un maduro de la base: las primeras
SEMILLA nt exactas (semilla + posición 1), a lo sumo MAX_MM diferencias en el
resto de la parte común, y una diferencia de largo de hasta HOLGURA nt (el 3'
varía por isomiRs). Se guarda la mejor coincidencia con sus diferencias
(y_ref, y_mm, y_dlen), así que medir la sensibilidad al umbral es releer la
tabla, no volver a comparar.

EXPORTAR solo saca los PRIMARIOS de los organismos de entrenamiento
(set_modelo en data/organismos.tsv). El duplicado no entra nunca: es la
validación independiente, y si entra deja de serlo.
"""
import argparse
import collections
import csv
import math
import os
import pathlib
import subprocess
import sys

RAIZ = pathlib.Path(__file__).resolve().parent.parent

# yasma tradeoff con los defaults de v1.1.1 escritos. Un default que cambie de
# versión no avisa: mismo criterio que trim.sh y align.sh.
TRADEOFF = [
    "--min_length", "15", "--max_length", "50",
    "--genome_scaling_factor", "0.4",
    "--merge_dist", "500", "--merge_strand_similarity", "0.5",
    "--max_skew", "0.9", "--skew_ignore_range", "20-24",
    "--min_complexity", "10", "--min_abundance", "50",
    "--min_abundance_density", "100", "--min_size", "19",
]
# tradeoff de v1.1.1 reserva 8 B por base y 8 B por base y por libreria antes
# de leer un read: -9 (sin memoria) en los cuatro proyectos que se probaron en
# Colab. El parche guarda lo mismo de forma dispersa y no cambia la salida
# (scripts/tradeoff_parche.py, tests/test_tradeoff_parche.py). Sin él no se anota.
PARCHE_TRADEOFF = [sys.executable, str(RAIZ / "scripts" / "tradeoff_parche.py"), "--verificar"]

# Columnas de loci.txt en v1.1.1 (assessClass.header). Si cambian, se corta:
# leer una columna corrida no falla, da features equivocadas.
CABECERA_LOCI = ["Locus", "Name", "Length", "Reads", "RPM", "UniqueReads", "FracTop",
                 "Strand", "MajorRNA", "MajorRNAReads", "Complexity", "Gap", "skew",
                 "size_1n", "size_1n_depth", "size_2n", "size_2n_depth",
                 "size_3n", "size_3n_depth", "sizecall"]

# Largos: miRNA ~21-22, siRNA vegetal 24, piRNA 26-31, tRF 30-40. Los bordes
# abiertos (15-17, 26-30, 31-50) juntan lo que en ningún reino es una clase.
BINS = [("15_17", 15, 17), ("18", 18, 18), ("19", 19, 19), ("20", 20, 20),
        ("21", 21, 21), ("22", 22, 22), ("23", 23, 23), ("24", 24, 24),
        ("25", 25, 25), ("26_30", 26, 30), ("31_50", 31, 50)]
# Ventana del pliegue a cada lado del RNA mayoritario. Un pre-miRNA animal mide
# ~60-70 nt y el maduro puede estar en cualquier brazo: con 80 de cada lado la
# horquilla entera entra sin importar el brazo. En plantas el precursor puede
# ser más largo; la horquilla local alrededor del maduro sigue entrando.
FLANCO = 80

SEMILLA = 8
MAX_MM = 1
HOLGURA = 3

FEATURES = (
    ["log_largo", "log_reads", "log_rpm", "log_unicos", "complejidad", "skew", "sesgo_hebra"]
    + [f"frac_{b}" for b, _, _ in BINS]
    + ["frac_5A", "frac_5C", "frac_5G", "frac_5U",
       "precision_5p", "frac_multimapeo", "frac_librerias",
       "largo_mayoritario", "gc_mayoritario", "mayoritario_5U",
       "mfe", "mfe_por_nt", "frac_pareado_mayoritario"]
)
META = ["org", "rol", "locus_id", "locus", "nombre", "hebra", "mayoritario", "sizecall",
        "reads_yasma", "reads_bam"]
ETIQUETA = ["y", "y_ref", "y_mm", "y_dlen"]

COMPLEMENTO = str.maketrans("ACGTN", "TGCAN")


def adn(seq):
    """Mayúsculas y U -> T: una sola convención para comparar."""
    return seq.upper().replace("U", "T")


def revcomp(seq):
    return seq.translate(COMPLEMENTO)[::-1]


# --------------------------------------------------------------------------
# Anotar
# --------------------------------------------------------------------------
def encontrar_loci(d):
    """El loci.txt de una anotación. Tiene que haber uno solo: con dos (otra
    corrida con -n distinto) no se adivina cuál."""
    d = pathlib.Path(d)
    hall = sorted(d.glob("tradeoff*/loci.txt")) + ([d / "loci.txt"] if (d / "loci.txt").exists() else [])
    if len(hall) != 1:
        raise SystemExit(f"esperaba un loci.txt en {d} (tradeoff*/loci.txt o loci.txt) "
                         f"y hay {len(hall)}: {', '.join(map(str, hall)) or 'ninguno'}")
    return hall[0]


def leer_loci(ruta):
    """Filas de loci.txt con reads. Corta si la cabecera no es la de v1.1.1."""
    with open(ruta) as h:
        r = csv.DictReader(h, delimiter="\t")
        if r.fieldnames != CABECERA_LOCI:
            raise SystemExit(f"{ruta}: la cabecera no es la de YASMA v1.1.1.\n"
                             f"  esperaba: {CABECERA_LOCI}\n  hay     : {r.fieldnames}")
        filas = [f for f in r if f["Reads"] not in ("0", "", "NA") and f["MajorRNA"] not in ("", "NA")]
    return filas


def anotar(bam, d, yasma="yasma", parche=None):
    """yasma tradeoff sobre el BAM, parado adentro de -o (la trampa de rutas de
    YASMA: guarda rutas relativas y las abre desde el CWD).

    Exige el parche de memoria (`parche` es el comando que lo verifica; los
    bancos lo reemplazan) e informa la memoria pico, que es lo que decide en
    qué máquina entra cada proyecto."""
    bam, d = pathlib.Path(bam).resolve(), pathlib.Path(d).resolve()
    if not bam.exists():
        raise SystemExit(f"no existe el BAM {bam}")
    r = subprocess.run(PARCHE_TRADEOFF if parche is None else parche,
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise SystemExit("falta el parche de memoria de yasma tradeoff: sin él, un genoma "
                         "de cientos de Mb con varias librerías muere con -9.\n"
                         "  Aplicalo con: ./scripts/tradeoff_parche.py\n" + (r.stdout + r.stderr).strip())
    d.mkdir(parents=True, exist_ok=True)
    cmd = [yasma, "tradeoff", "-a", str(bam), "-o", str(d), *TRADEOFF]
    print("$ " + " ".join(cmd), flush=True)
    p = subprocess.Popen(cmd, cwd=d)
    _, st, ru = os.wait4(p.pid, 0)
    rc, pico = os.waitstatus_to_exitcode(st), ru.ru_maxrss * 1024
    if rc in (-9, 137):
        raise SystemExit(f"yasma tradeoff se quedó sin memoria (SIGKILL, rc={rc}) tras llegar a "
                         f"{pico/1e9:.1f} GB: este proyecto necesita una máquina con más RAM")
    if rc != 0:
        raise SystemExit(f"yasma tradeoff salió con {rc}: mirá {d}/tradeoff/log.txt")
    print(f"memoria pico de tradeoff: {pico/1e9:.2f} GB", flush=True)
    loci = encontrar_loci(d)
    n = len(leer_loci(loci))
    if n == 0:
        raise SystemExit(f"{loci}: yasma terminó bien pero no anotó ningún locus con reads")
    print(f"{n} loci anotados -> {loci}")
    return loci


# --------------------------------------------------------------------------
# Features
# --------------------------------------------------------------------------
class Locus:
    """Lo que hace falta juntar de los reads de un locus, en una sola pasada."""

    def __init__(self, mayoritario):
        self.mayoritario = adn(mayoritario)
        self.n = 0
        self.largos = collections.Counter()
        self.nt5 = collections.Counter()
        self.pos5 = {"+": collections.Counter(), "-": collections.Counter()}
        self.xy = collections.Counter()
        self.rgs = set()
        self.pos_may = collections.Counter()

    def agregar(self, seq, hebra, ini, fin, rg, xy):
        seq = adn(seq)
        self.n += 1
        self.largos[len(seq)] += 1
        if seq:
            self.nt5[seq[0]] += 1
        # El 5' de un read en la hebra - es su último nt en el genoma.
        self.pos5[hebra][ini if hebra == "+" else fin - 1] += 1
        if xy:
            self.xy[xy] += 1
        if rg is not None:
            self.rgs.add(rg)
        if seq == self.mayoritario:
            self.pos_may[(hebra, ini, fin)] += 1


def plegar(loc, chrom, genoma, largo_chrom, fold, flanco=FLANCO):
    """Pliega una ventana de ±flanco alrededor del RNA mayoritario, en su hebra.
    Sin el RNA mayoritario entre los reads del locus no hay dónde centrar la
    ventana: NaN, y exportar se niega a usar esa fila."""
    nan = float("nan")
    if not loc.pos_may:
        return {"mfe": nan, "mfe_por_nt": nan, "frac_pareado_mayoritario": nan}
    (hebra, ini, fin), _ = loc.pos_may.most_common(1)[0]
    w_ini, w_fin = max(0, ini - flanco), min(largo_chrom, fin + flanco)
    seq = adn(genoma(chrom, w_ini, w_fin))
    off = ini - w_ini
    if hebra == "-":
        seq, off = revcomp(seq), w_fin - fin
    estructura, mfe = fold(seq.replace("T", "U"))
    largo = fin - ini
    tramo = estructura[off:off + largo]
    return {"mfe": round(float(mfe), 3),
            "mfe_por_nt": round(float(mfe) / len(seq), 5),
            "frac_pareado_mayoritario": round(sum(c in "()" for c in tramo) / largo, 4)}


def features_de(fila, loc, n_libs, pliegue):
    reads = int(fila["Reads"])
    f = {
        "log_largo": math.log10(int(fila["Length"]) + 1),
        "log_reads": math.log10(reads + 1),
        "log_rpm": math.log10(float(fila["RPM"]) + 1),
        "log_unicos": math.log10(int(fila["UniqueReads"]) + 1),
        "complejidad": float(fila["Complexity"]),
        "skew": float(fila["skew"]),
        "sesgo_hebra": abs(2 * float(fila["FracTop"]) - 1),
    }
    n = loc.n or 1
    for b, lo, hi in BINS:
        f[f"frac_{b}"] = sum(c for largo, c in loc.largos.items() if lo <= largo <= hi) / n
    for nt, nombre in (("A", "A"), ("C", "C"), ("G", "G"), ("T", "U")):
        f[f"frac_5{nombre}"] = loc.nt5[nt] / n
    dom = "+" if sum(loc.pos5["+"].values()) >= sum(loc.pos5["-"].values()) else "-"
    tot = sum(loc.pos5[dom].values())
    f["precision_5p"] = max(loc.pos5[dom].values()) / tot if tot else 0.0
    colocados = loc.xy["U"] + loc.xy["P"] + loc.xy["R"]
    f["frac_multimapeo"] = (loc.xy["P"] + loc.xy["R"]) / colocados if colocados else 0.0
    f["frac_librerias"] = len(loc.rgs) / n_libs if n_libs else 0.0
    may = loc.mayoritario
    f["largo_mayoritario"] = len(may)
    f["gc_mayoritario"] = (may.count("G") + may.count("C")) / len(may) if may else 0.0
    f["mayoritario_5U"] = 1 if may.startswith("T") else 0
    f.update(pliegue)
    return {k: (round(v, 6) if isinstance(v, float) and not math.isnan(v) else v) for k, v in f.items()}


def partir_locus(locus):
    chrom, rango = locus.rsplit(":", 1)
    ini, fin = rango.split("-")
    return chrom, int(ini), int(fin)


def construir(filas, reads_de, n_libs, genoma, largos_chrom, fold, org, rol, flanco=FLANCO):
    """Una fila de matriz por locus. reads_de(chrom, ini, fin) da tuplas
    (seq, hebra, ini, fin, rg, xy) de los reads mapeados que solapan."""
    salida = []
    for i, fila in enumerate(filas, 1):
        chrom, ini, fin = partir_locus(fila["Locus"])
        loc = Locus(fila["MajorRNA"])
        for r in reads_de(chrom, ini, fin):
            loc.agregar(*r)
        pliegue = plegar(loc, chrom, genoma, largos_chrom[chrom], fold, flanco)
        fila_m = {
            "org": org, "rol": rol, "locus_id": f"{org}_{rol}:{fila['Name']}",
            "locus": fila["Locus"], "nombre": fila["Name"], "hebra": fila["Strand"],
            "mayoritario": fila["MajorRNA"], "sizecall": fila["sizecall"],
            "reads_yasma": int(fila["Reads"]), "reads_bam": loc.n,
        }
        fila_m.update(features_de(fila, loc, n_libs, pliegue))
        salida.append(fila_m)
        if i % 2000 == 0:
            print(f"   {i}/{len(filas)} loci", flush=True)
    return salida


def escribir(filas, ruta, columnas):
    with open(ruta, "w", newline="") as h:
        w = csv.DictWriter(h, fieldnames=columnas, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        w.writerows(filas)


def leer_tsv(ruta):
    with open(ruta) as h:
        r = csv.DictReader(h, delimiter="\t")
        return r.fieldnames, list(r)


# --- adaptadores a pysam y ViennaRNA: lo único que el banco no puede fingir ---
def reads_pysam(bamf):
    def reads_de(chrom, ini, fin):
        for r in bamf.fetch(contig=chrom, start=ini, stop=fin):
            if r.is_unmapped:
                continue
            yield (r.get_forward_sequence(), "-" if r.is_reverse else "+",
                   r.reference_start, r.reference_end,
                   r.get_tag("RG") if r.has_tag("RG") else None,
                   r.get_tag("XY") if r.has_tag("XY") else None)
    return reads_de


def fold_vienna():
    import RNA
    return lambda seq: tuple(RNA.fold(seq))


def cmd_features(a):
    import pysam
    filas = leer_loci(encontrar_loci(a.dir))
    bamf = pysam.AlignmentFile(str(a.bam), "rb")
    if not bamf.has_index():
        raise SystemExit(f"{a.bam} no tiene índice (.bai): samtools index {a.bam}")
    n_libs = len(bamf.header.to_dict().get("RG", []))
    fasta = pysam.FastaFile(str(a.genoma))
    faltan = sorted(set(bamf.references) - set(fasta.references))
    if faltan:
        raise SystemExit(f"el genoma {a.genoma} no tiene {len(faltan)} secuencias del BAM "
                         f"({', '.join(faltan[:3])}...): no es el FASTA contra el que se alineó")
    largos = dict(zip(fasta.references, fasta.lengths))
    print(f"{len(filas)} loci, {n_libs} librerías, genoma {a.genoma}")
    m = construir(filas, reads_pysam(bamf), n_libs, fasta.fetch, largos, fold_vienna(),
                  a.org, a.rol, a.flanco)
    escribir(m, a.o, META + FEATURES)
    distintos = sum(1 for f in m if f["reads_yasma"] != f["reads_bam"])
    print(f"{len(m)} filas -> {a.o}")
    if distintos:
        print(f"AVISO: en {distintos} loci los reads contados acá no coinciden con los de "
              f"loci.txt. La regla de conteo es la de YASMA; si difiere, no se está mirando "
              f"lo mismo que YASMA miró (¿otro BAM?).")
    return 1 if distintos else 0


# --------------------------------------------------------------------------
# Etiquetar
# --------------------------------------------------------------------------
def leer_fasta(rutas):
    seqs, nombre = [], None
    for ruta in rutas:
        with open(ruta) as h:
            for linea in h:
                linea = linea.strip()
                if linea.startswith(">"):
                    nombre = linea[1:].split()[0]
                elif linea and nombre:
                    seqs.append((nombre, adn(linea)))
                    nombre = None
    return seqs


def indice_maduros(maduros, semilla=SEMILLA):
    ind = collections.defaultdict(list)
    for nombre, seq in maduros:
        if len(seq) >= semilla:
            ind[seq[:semilla]].append((nombre, seq))
    return ind


def mejor_coincidencia(seq, indice, semilla=SEMILLA):
    """(ref, mm, dlen) del maduro más parecido que comparte la semilla exacta, o
    None. mm cuenta diferencias en la parte común después de la semilla; dlen es
    largo(locus) - largo(ref). Sin desplazar el 5': un corrimiento de una base
    ya es otro miRNA."""
    seq = adn(seq)
    mejor, clave = None, None
    for nombre, ref in indice.get(seq[:semilla], []):
        comun = min(len(seq), len(ref))
        mm = sum(1 for i in range(semilla, comun) if seq[i] != ref[i])
        dlen = len(seq) - len(ref)
        k = (mm, abs(dlen), nombre)
        if clave is None or k < clave:
            mejor, clave = (nombre, mm, dlen), k
    return mejor


def etiquetar(filas, maduros, semilla=SEMILLA, max_mm=MAX_MM, holgura=HOLGURA):
    ind = indice_maduros(maduros, semilla)
    for f in filas:
        m = mejor_coincidencia(f["mayoritario"], ind, semilla)
        if m is None:
            f.update(y=0, y_ref="", y_mm="", y_dlen="")
        else:
            ref, mm, dlen = m
            f.update(y=int(mm <= max_mm and abs(dlen) <= holgura), y_ref=ref, y_mm=mm, y_dlen=dlen)
    return filas


def cmd_etiquetar(a):
    cols, filas = leer_tsv(a.matriz)
    maduros = leer_fasta(a.maduros)
    if not maduros:
        raise SystemExit(f"no leí ninguna secuencia de {', '.join(map(str, a.maduros))}")
    etiquetar(filas, maduros, a.semilla, a.max_mm, a.holgura)
    cols = [c for c in cols if c not in ETIQUETA] + ETIQUETA
    escribir(filas, a.o or a.matriz, cols)
    pos = sum(int(f["y"]) for f in filas)
    cerca = sum(1 for f in filas if f["y_ref"] and not int(f["y"]))
    print(f"{len(maduros)} maduros; {pos} de {len(filas)} loci positivos "
          f"(semilla {a.semilla}, ≤{a.max_mm} mm, ±{a.holgura} nt); "
          f"{cerca} comparten semilla pero quedan afuera del umbral")
    return 0


# --------------------------------------------------------------------------
# Exportar para loo_cv.py
# --------------------------------------------------------------------------
def set_modelo(organismos=RAIZ / "data" / "organismos.tsv"):
    with open(organismos) as h:
        lineas = [l for l in h if not l.startswith("#")]
    return {f["org"]: f["set_modelo"] for f in csv.DictReader(lineas, delimiter="\t")}


def exportar(rutas, organismos=RAIZ / "data" / "organismos.tsv"):
    """Filas (org, locus_id, y, FEATURES) de los primarios de entrenamiento, y un
    resumen de lo que quedó afuera y por qué."""
    sm = set_modelo(organismos)
    salida, afuera, vistos = [], collections.Counter(), set()
    for ruta in rutas:
        cols, filas = leer_tsv(ruta)
        faltan = [c for c in ["org", "rol", "locus_id", "y"] + FEATURES if c not in cols]
        if faltan:
            raise SystemExit(f"{ruta}: faltan columnas {faltan}"
                             + (" — ¿se etiquetó? (matriz.py etiquetar)" if "y" in faltan else ""))
        for f in filas:
            org, rol = f["org"], f["rol"]
            if sm.get(org) != "entrenamiento":
                afuera[(org, rol, "aplicación: se predice, no se entrena")] += 1
                continue
            if rol != "primario":
                afuera[(org, rol, "duplicado: es la validación independiente")] += 1
                continue
            if f["y"] not in ("0", "1"):
                raise SystemExit(f"{ruta}: {f['locus_id']} tiene y={f['y']!r}; tiene que ser 0 o 1")
            malas = [c for c in FEATURES if f[c] in ("", "nan", "NaN") or not math.isfinite(float(f[c]))]
            if malas:
                raise SystemExit(f"{ruta}: {f['locus_id']} tiene features vacías o NaN: {malas}")
            if f["locus_id"] in vistos:
                raise SystemExit(f"locus_id repetido: {f['locus_id']} (¿la misma matriz dos veces?)")
            vistos.add(f["locus_id"])
            salida.append({"org": org, "locus_id": f["locus_id"], "y": f["y"],
                           **{c: f[c] for c in FEATURES}})
    return salida, afuera


def cmd_exportar(a):
    filas, afuera = exportar(a.matrices)
    for (org, rol, por), n in sorted(afuera.items()):
        print(f"   afuera {org}/{rol}: {n} loci ({por})")
    if not filas:
        raise SystemExit("no quedó ninguna fila de entrenamiento: hacen falta primarios de "
                         "gadmo, galga o maggi, etiquetados")
    escribir(filas, a.o, ["org", "locus_id", "y"] + FEATURES)
    por_org = collections.Counter((f["org"], f["y"]) for f in filas)
    for org in sorted({o for o, _ in por_org}):
        print(f"   {org}: {por_org[(org, '1')]} positivos, {por_org[(org, '0')]} sin etiquetar")
    print(f"{len(filas)} filas, {len(FEATURES)} features -> {a.o}")
    if len({f['org'] for f in filas}) < 3:
        print("AVISO: loo_cv.py deja un organismo afuera por vez; con menos de 3 organismos "
              "de entrenamiento todavía no mide transferencia.")
    return 0


# --------------------------------------------------------------------------
# Lo que usa el notebook (30_matriz): estado en Drive, maduros, registro
# --------------------------------------------------------------------------
def estado_drive(drive, organismos=RAIZ / "data" / "organismos.tsv"):
    """Qué proyectos tienen BAM en Drive (10_bam/<org>/<rol>.bam: solo llega lo
    que pasó align.sh verificar) y cuáles ya tienen anotación y matriz."""
    drive, sm = pathlib.Path(drive), set_modelo(organismos)
    filas = []
    for bam in sorted(drive.glob("10_bam/*/*.bam")):
        org, rol = bam.parent.name, bam.name[:-len(".bam")]
        filas.append({
            "org": org, "rol": rol, "set_modelo": sm.get(org, "?"),
            "bam_gb": round(bam.stat().st_size / 1e9, 2),
            "anotacion": (drive / "20_yasma" / org / rol / "tradeoff" / "loci.txt").exists(),
            "matriz": (drive / "40_features" / org / f"{rol}.matriz.tsv").exists(),
        })
    return filas


def validar_maduros(texto, minimo=20):
    """(n, None) si es un FASTA de maduros de verdad, (n, motivo) si no. Una URL
    que cambió devuelve una página HTML con código 200: sin esto se etiquetaría
    contra nada y TODO saldría sin etiquetar, sin ningún error."""
    lineas = [l.strip() for l in texto.splitlines() if l.strip()]
    if not lineas or not lineas[0].startswith(">"):
        return 0, "no empieza con '>': no es un FASTA (¿una página de error?)"
    seqs = [l for l in lineas if not l.startswith(">")]
    raras = [q for q in seqs if not (17 <= len(q) <= 28 and set(adn(q)) <= set("ACGT"))]
    if raras:
        return len(seqs), (f"{len(raras)} secuencias que no parecen maduros (17-28 nt de ACGU), "
                           f"p. ej. {raras[0][:40]!r}: ¿precursores en vez de maduros?")
    if len(seqs) < minimo:
        return len(seqs), f"solo {len(seqs)} secuencias: esperaba al menos {minimo}"
    return len(seqs), None


def resumen_matriz(ruta):
    """Lo que va a data/matrices.tsv: cuántos loci, cuántos positivos, y el
    sha256 del fichero, que es el ancla de qué matriz se usó."""
    import hashlib
    cols, filas = leer_tsv(ruta)
    y = [f["y"] for f in filas] if "y" in cols else []
    return {"loci": len(filas),
            "positivos": sum(v == "1" for v in y) if y else "",
            "features": sum(c in cols for c in FEATURES),
            "sha256": hashlib.sha256(pathlib.Path(ruta).read_bytes()).hexdigest()}


# --------------------------------------------------------------------------
# Excel: para mirar las matrices, no para entrenar
# --------------------------------------------------------------------------
# Qué es cada columna, de dónde sale y por qué está. Es la tabla de
# docs/matriz.md, para que el Excel se lea solo. El banco exige que cada
# feature tenga su fila: una columna sin explicar es un número sin unidad.
_LARGOS = {"15_17": "15-17", "26_30": "26-30", "31_50": "31-50"}
DICCIONARIO = (
    [("identificación", c, d) for c, d in (
        ("org", "Organismo (código del proyecto)."),
        ("rol", "primario = descubrimiento; duplicado = validación independiente. "
                "El duplicado nunca entra al entrenamiento."),
        ("locus_id", "Identificador único: <org>_<rol>:<nombre del locus>."),
        ("locus", "Coordenadas chrom:inicio-fin, como las da YASMA."),
        ("nombre", "Nombre del locus en loci.txt."),
        ("hebra", "Hebra del locus según YASMA (+, - o . si es mixta)."),
        ("mayoritario", "RNA mayoritario del locus (la secuencia más abundante)."),
        ("sizecall", "Clase de tamaño que llama YASMA (15-30 nt); N si no hay una dominante."),
        ("reads_yasma", "Reads del locus según loci.txt."),
        ("reads_bam", "Reads contados acá con la regla de YASMA. Tiene que ser igual a reads_yasma."),
    )]
    + [("etiqueta", c, d) for c, d in (
        ("y", "1 = el mayoritario coincide con un maduro de la base (positivo). "
              "0 = SIN ETIQUETAR, no negativo."),
        ("y_ref", "El maduro más parecido, aunque quede fuera del umbral."),
        ("y_mm", "Diferencias con ese maduro fuera de la semilla."),
        ("y_dlen", "Diferencia de largo con ese maduro (isomiRs del 3')."),
    )]
    + [("feature", c, d) for c, d in (
        ("log_largo", "log10(largo del locus + 1). Los loci de miRNA son cortos; los de siRNA, largos."),
        ("log_reads", "log10(reads + 1): abundancia."),
        ("log_rpm", "log10(RPM + 1): abundancia normalizada por profundidad."),
        ("log_unicos", "log10(secuencias únicas + 1). Un miRNA tiene pocas; un siRNA, muchas."),
        ("complejidad", "Secuencias únicas por nt del locus (Complexity de YASMA)."),
        ("skew", "Fracción del read más abundante. Alta en miRNAs, que se procesan con precisión."),
        ("sesgo_hebra", "|2·FracTop − 1|: 1 si todo está en una hebra, 0 si está repartido."),
    )]
    + [("feature", f"frac_{b}", f"Fracción de reads de {_LARGOS.get(b, b)} nt.") for b, _, _ in BINS]
    + [("feature", c, d) for c, d in (
        ("frac_5A", "Fracción de reads con A en el 5'."),
        ("frac_5C", "Fracción de reads con C en el 5'."),
        ("frac_5G", "Fracción de reads con G en el 5'."),
        ("frac_5U", "Fracción de reads con U en el 5'. La U 5' la prefiere AGO1 (miRNAs)."),
        ("precision_5p", "En la hebra dominante, fracción de reads que empieza en la posición 5' "
                         "más común. Dicer corta preciso; la degradación no."),
        ("frac_multimapeo", "Reads colocados por sorteo (XY P o R) sobre los colocados (U+P+R). "
                            "Alta en repetidos y transposones."),
        ("frac_librerias", "Fracción de las librerías del proyecto donde aparece el locus."),
        ("largo_mayoritario", "Largo del RNA mayoritario (nt)."),
        ("gc_mayoritario", "Fracción GC del RNA mayoritario."),
        ("mayoritario_5U", "1 si el RNA mayoritario empieza con U."),
        ("mfe", "Energía mínima del pliegue (kcal/mol) de ±80 nt alrededor del mayoritario, "
                "en su hebra. Más negativo = más estable."),
        ("mfe_por_nt", "mfe dividida por el largo de la ventana. Un precursor de miRNA pliega "
                       "en horquilla estable."),
        ("frac_pareado_mayoritario", "Fracción del mayoritario apareada en ese pliegue. "
                                     "Un miRNA maduro vive en el tallo."),
    )]
)
# Enteros en el Excel; el resto de las numéricas con 3 decimales.
_ENTEROS = {"reads_yasma", "reads_bam", "y", "y_mm", "y_dlen", "largo_mayoritario",
            "mayoritario_5U"}
_TEXTO = {"org", "rol", "locus_id", "locus", "nombre", "hebra", "mayoritario", "sizecall",
          "y_ref"}
FILAS_MAX_EXCEL = 1_048_575


def _valor(col, v):
    """El valor de la celda: número si la columna es numérica, vacío si NaN."""
    if col in _TEXTO:
        return v
    if v is None or v.strip() == "" or v.strip().lower() == "nan":
        return None
    try:
        x = float(v)
    except ValueError:
        return v
    if math.isnan(x):
        return None
    return int(x) if col in _ENTEROS and x.is_integer() else x


def excel(rutas, salida, organismos=RAIZ / "data" / "organismos.tsv"):
    """Un libro con Resumen, Diccionario y una hoja por proyecto. -> hojas."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    letra = Font(name="Arial", size=10)
    negrita = Font(name="Arial", size=10, bold=True)
    gris = PatternFill("solid", start_color="D9D9D9")
    sm = set_modelo(organismos)

    proyectos = []
    for r in rutas:
        cols, filas = leer_tsv(r)
        faltan = [c for c in META + FEATURES if c not in (cols or [])]
        if faltan:
            raise SystemExit(f"{r}: no es una matriz de matriz.py (faltan {faltan[:3]}...)")
        if len(filas) > FILAS_MAX_EXCEL:
            raise SystemExit(f"{r}: {len(filas)} loci no entran en una hoja de Excel")
        org = filas[0]["org"] if filas else pathlib.Path(r).parent.name
        rol = filas[0]["rol"] if filas else pathlib.Path(r).name.split(".")[0]
        hoja = f"{org}_{rol}"[:31]
        if hoja in {p[0] for p in proyectos}:
            raise SystemExit(f"dos matrices de {org}/{rol}: pasá una sola")
        orden = META + [c for c in ETIQUETA if c in cols] + FEATURES
        proyectos.append((hoja, org, rol, orden, filas))

    wb = Workbook(write_only=True)
    # Las fórmulas del Resumen se escriben sin valor calculado: que Excel y
    # Sheets las calculen al abrir, en vez de mostrar celdas vacías.
    wb.calculation.fullCalcOnLoad = True

    def celda(ws, v, font=letra, fmt=None, fill=None):
        c = WriteOnlyCell(ws, value=v)
        c.font = font
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        return c

    # Resumen: los conteos son fórmulas sobre las hojas, no números copiados.
    ws = wb.create_sheet("Resumen")
    for letra_col, ancho in zip("ABCDEFG", (20, 10, 11, 14, 9, 14, 20)):
        ws.column_dimensions[letra_col].width = ancho
    ws.freeze_panes = "A2"
    ws.append([celda(ws, h, negrita, fill=gris) for h in
               ("proyecto", "org", "rol", "set_modelo", "loci", "positivos (y=1)", "hoja")])
    for hoja, org, rol, orden, _ in proyectos:
        loci = f"=COUNTA('{hoja}'!C:C)-1"
        if "y" in orden:
            col_y = _letra(orden.index("y") + 1)
            pos = f"=COUNTIF('{hoja}'!{col_y}:{col_y},1)"
        else:
            pos = "sin etiquetar"
        ws.append([celda(ws, v) for v in (f"{org}/{rol}", org, rol, sm.get(org, "?"))]
                  + [celda(ws, loci, fmt="#,##0"), celda(ws, pos, fmt="#,##0"), celda(ws, hoja)])
    ws.append([])
    ws.append([celda(ws, "y = 0 es SIN ETIQUETAR, no negativo (PU learning). "
                         "Qué es cada columna: hoja Diccionario. Fuente: scripts/matriz.py, "
                         "docs/matriz.md.")])

    ws = wb.create_sheet("Diccionario")
    for letra_col, ancho in zip("ABC", (16, 26, 100)):
        ws.column_dimensions[letra_col].width = ancho
    ws.freeze_panes = "A2"
    ws.append([celda(ws, h, negrita, fill=gris) for h in ("grupo", "columna", "qué es")])
    for grupo, col, desc in DICCIONARIO:
        c = celda(ws, desc)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([celda(ws, grupo), celda(ws, col), c])

    for hoja, _org, _rol, orden, filas in proyectos:
        ws = wb.create_sheet(hoja)
        for i, col in enumerate(orden, 1):
            ws.column_dimensions[_letra(i)].width = (
                30 if col in ("locus_id", "mayoritario", "y_ref") else max(10, len(col) + 2))
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = f"A1:{_letra(len(orden))}{len(filas) + 1}"
        ws.append([celda(ws, c, negrita, fill=gris) for c in orden])
        for f in filas:
            ws.append([celda(ws, _valor(c, f.get(c)),
                             fmt=None if c in _TEXTO else ("0" if c in _ENTEROS else "0.000"))
                       for c in orden])
    wb.save(salida)
    return [p[0] for p in proyectos]


def _letra(n):
    """1 -> A, 27 -> AA."""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def cmd_excel(a):
    hojas = excel(a.matrices, a.o)
    print(f"{len(hojas)} proyecto(s) -> {a.o}: Resumen, Diccionario, {', '.join(hojas)}")
    return 0


# --------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("anotar", help="yasma tradeoff sobre el BAM")
    p.add_argument("bam", type=pathlib.Path)
    p.add_argument("dir", type=pathlib.Path)
    p.add_argument("--yasma", default=os.environ.get("YASMA", "yasma"))

    p = sub.add_parser("features", help="la matriz de un proyecto")
    p.add_argument("dir", type=pathlib.Path, help="directorio de la anotación")
    p.add_argument("bam", type=pathlib.Path)
    p.add_argument("--org", required=True)
    p.add_argument("--rol", required=True, choices=("primario", "duplicado"))
    p.add_argument("--genoma", required=True, type=pathlib.Path,
                   help="el FASTA descomprimido (align.sh fasta <org>)")
    p.add_argument("--flanco", type=int, default=FLANCO)
    p.add_argument("-o", required=True, type=pathlib.Path)

    p = sub.add_parser("etiquetar", help="y por secuencia contra maduros")
    p.add_argument("matriz", type=pathlib.Path)
    p.add_argument("maduros", type=pathlib.Path, nargs="+")
    p.add_argument("--semilla", type=int, default=SEMILLA)
    p.add_argument("--max-mm", type=int, default=MAX_MM)
    p.add_argument("--holgura", type=int, default=HOLGURA)
    p.add_argument("-o", type=pathlib.Path)

    p = sub.add_parser("exportar", help="features.tsv para loo_cv.py")
    p.add_argument("matrices", type=pathlib.Path, nargs="+")
    p.add_argument("-o", required=True, type=pathlib.Path)

    p = sub.add_parser("excel", help="las matrices en un .xlsx, con diccionario")
    p.add_argument("matrices", type=pathlib.Path, nargs="+")
    p.add_argument("-o", required=True, type=pathlib.Path)

    a = ap.parse_args(argv)
    if a.cmd == "anotar":
        anotar(a.bam, a.dir, a.yasma)
        return 0
    return {"features": cmd_features, "etiquetar": cmd_etiquetar, "exportar": cmd_exportar,
            "excel": cmd_excel}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main())
